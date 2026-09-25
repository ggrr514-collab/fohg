# -*- coding: utf-8 -*-
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

OUT = "/home/user/fohg/単元_素顔同盟_私/教員用評価シート.xlsx"
NAVY=PatternFill('solid',fgColor='1F3864'); LBLUE=PatternFill('solid',fgColor='E8F0FE')
ORANGE_F=PatternFill('solid',fgColor='FFF2CC'); PURPLE_F=PatternFill('solid',fgColor='E8DAEF')
GREEN_F=PatternFill('solid',fgColor='D5F5E3'); GRAY=PatternFill('solid',fgColor='F2F2F2')
INPUT=PatternFill('solid',fgColor='FFFFF0'); PINK=PatternFill('solid',fgColor='FADBD8')
wrap=Alignment(wrap_text=True,vertical='top'); wrapc=Alignment(wrap_text=True,vertical='center')
center=Alignment(horizontal='center',vertical='center',wrap_text=True)
thin=Side(style='thin',color='BFBFBF'); box=Border(left=thin,right=thin,top=thin,bottom=thin)
F_TITLE=Font(size=14,bold=True,color='FFFFFF'); F_HEAD=Font(size=10,bold=True,color='FFFFFF')
F_LBL=Font(size=10,bold=True,color='333333'); F_BODY=Font(size=10)
F_RED=Font(size=10,bold=True,color='C00000'); F_GOAL=Font(size=11,bold=True,color='1F3864')
F_SMALL=Font(size=9,color='555555')

GOAL = ("自分で決めた観点から二つの作品を批判的に読みながら、三つの論題について、"
        "文章に表れているものの見方や考え方について考えることができる。")
CRIT_S = ("「読むこと」において、文章を批判的に読みながら、"
          "文章に表れているものの見方や考え方について考えている。")
CRIT_A = ("積極的に文章を批判的に読み、学習課題に沿って、"
          "考えたことを互いに伝え合おうとしている。")

def mg(ws,r1,c1,r2,c2,val,font=None,fill=None,align=None,border=False,h=None):
    ws.cell(row=r1,column=c1).value=val
    if r1!=r2 or c1!=c2: ws.merge_cells(start_row=r1,start_column=c1,end_row=r2,end_column=c2)
    for r in range(r1,r2+1):
        for c in range(c1,c2+1):
            cell=ws.cell(row=r,column=c)
            if font: cell.font=font
            if fill: cell.fill=fill
            if align: cell.alignment=align
            if border: cell.border=box
    if h: ws.row_dimensions[r1].height=h

wb=openpyxl.Workbook()

# ============ Sheet1 回答入力（14列） ============
ws=wb.active; ws.title='回答入力'
cols=[('クラス',8),('出席番号',9),('氏名',14),('選んだ観点',18),
      ('論題1 仮面とデータは同じものか',30),
      ('論題2 訴えた人は、救われたか',30),
      ('論題3 続けさせているのは、誰か',30),
      ('聞いて考えた①',24),('聞いて考えた②',24),('聞いて考えた③',24),
      ('AI評価\n（記述のみ暫定）',12),('AIコメント',30),
      ('★教師最終評価\n（記述＋発言）',13),('ふりかえり',24)]
for i,(name,w) in enumerate(cols,start=1):
    ws.column_dimensions[get_column_letter(i)].width=w
mg(ws,1,1,1,14,'教員用評価シート　「私」を消すもの、「私」を支えるもの（中3・全7時間／まとめ＝グループディスカッション）',
   F_TITLE,NAVY,center,h=28)
mg(ws,2,1,2,14,'目標：'+GOAL,F_GOAL,ORANGE_F,wrapc,h=32)
mg(ws,3,1,3,14,'★ 思考・判断・表現は【D〜J列の記述】と【当日の発言（「協議見取り」シート）】を合わせてM列で確定する。'
                'K列のAI評価は記述だけを見た暫定値。生徒シートへの自動返却は行わない（評価は口頭または紙で返す）。',
   F_RED,PINK,wrapc,h=30)
fills={11:PURPLE_F,12:PURPLE_F,13:PINK,14:GREEN_F}
for i,(name,_) in enumerate(cols,start=1):
    c=ws.cell(row=4,column=i); c.value=name
    c.font=F_HEAD; c.fill=NAVY; c.alignment=center; c.border=box
ws.row_dimensions[4].height=42
dv_ai=DataValidation(type='list',formula1='"S,A,B,C,D"',allow_blank=True,showDropDown=False)
dv_t=DataValidation(type='list',formula1='"S,A,B,C,D"',allow_blank=True,showDropDown=False)
ws.add_data_validation(dv_ai); ws.add_data_validation(dv_t)
for r in range(5,45):
    for i in range(1,15):
        c=ws.cell(row=r,column=i); c.border=box; c.font=F_BODY
        c.alignment=center if i in (1,2,11,13) else wrap
        if i in fills: c.fill=fills[i]
    ws.row_dimensions[r].height=30
dv_ai.add('K5:K44'); dv_t.add('M5:M44')
ws.freeze_panes='D5'

# ============ Sheet2 AI設定 ============
ws=wb.create_sheet('AI設定')
ws.column_dimensions['A'].width=26; ws.column_dimensions['B'].width=58
mg(ws,1,1,1,2,'AI設定',F_TITLE,NAVY,center,h=28)
for r,a,b in [(3,'AIモデル','gemini-2.0-flash'),(4,'Temperature',0.1),
              (6,'評価対象','論題1〜3（E〜G列）＋聞いて考えたこと（H〜J列）'),
              (7,'評価観点','思考・判断・表現のみ　5段階（S/A/B/C/D）'),
              (8,'知識・技能','この単元では評価しない（指導のみ。別単元・定期テストで評価）'),
              (10,'データ範囲','5行目〜44行目'),
              (11,'ClassroomフォルダID','（ここにフォルダIDを入力）')]:
    mg(ws,r,1,r,1,a,F_LBL,GRAY,wrapc,border=True)
    mg(ws,r,2,r,2,b,F_BODY,INPUT,wrapc,border=True,h=24)
mg(ws,13,1,13,2,'★ APIキーはこのシートに書かない。GASの「スクリプト プロパティ」に GEMINI_API_KEY として登録する。',
   F_RED,PINK,wrap,h=28)
mg(ws,15,1,15,2,'★ B/Aの分水嶺：条件①（3つの論題すべてを観点と本文の根拠でまとめた）と'
                 '条件②（当日、自分の観点から根拠を示して発言した）の両方がなければA以上にしない。',
   F_RED,None,wrap,h=44)
mg(ws,17,1,17,2,'★ AIは条件②（当日の発言）を判定できない。K列は記述だけを見た暫定評価。'
                 '「協議見取り」H列（発言の判定）と合わせて、M列で最終評価を確定する。',
   F_RED,None,wrap,h=44)
mg(ws,19,1,19,2,'★ 生徒シートに評価を書き戻す機能は無い（生徒シートの「AIの評価と自分の評価」欄を'
                 '廃止したため）。評価は口頭または紙で返す。',F_RED,None,wrap,h=40)
mg(ws,21,1,21,1,'評価規準（思判表）',F_LBL,GRAY,wrapc,border=True)
mg(ws,21,2,21,2,CRIT_S,F_BODY,LBLUE,wrapc,border=True,h=40)
mg(ws,22,1,22,1,'評価規準（態度）',F_LBL,GRAY,wrapc,border=True)
mg(ws,22,2,22,2,CRIT_A,F_BODY,LBLUE,wrapc,border=True,h=34)
mg(ws,24,1,24,2,'★ 観点別評価への読みかえ（指導要録はA・B・Cの3段階）　'
                 'S・A → A（十分満足できる）／B → B（おおむね満足できる）／C・D → C（努力を要する）。'
                 '生徒への返却はS〜Dのまま。',F_RED,PINK,wrap,h=44)

# ============ Sheet3 協議見取り（10列） ============
ws=wb.create_sheet('協議見取り')
widths=[8,9,14,7,22,8,28,10,28,10]
for i,w in enumerate(widths,start=1): ws.column_dimensions[get_column_letter(i)].width=w
mg(ws,1,1,1,10,'協議見取り記録（2時・5〜7時）　―― 当日の発言を評価する表。思判表の条件②と、態度の評価材',
   F_TITLE,NAVY,center,h=28)
mg(ws,2,1,2,10,'★ 発話は残らない。ここを書かないと条件②（当日の発言）と態度の評価材が無くなる。'
                '5・6時は巡回しながらF〜J列を埋める。全発言を聞き取る必要はない。',F_RED,PINK,wrapc,h=32)
mg(ws,3,1,3,10,'★ E列は2時のおわりに記入し、【観点がばらけるように】班を組む。観点は9つから選ばせる（必修は「批評・評価」）。'
                '論題は3つ。1=仮面とデータは同じものか／2=訴えた人は救われたか／3=続けさせているのは誰か。'
                '発表する論題は当日に知らせる（生徒は3つ全部準備している）。',F_BODY,ORANGE_F,wrapc,h=40)
heads=['クラス','出席番号','氏名','班','選んだ観点（2時）','発表論題',
       '当日の発言の要点','条件②\n○△×','ほかの観点を受けて述べ直した箇所','態度\nA/B/C']
for i,h in enumerate(heads,start=1):
    c=ws.cell(row=4,column=i); c.value=h
    c.font=F_HEAD; c.fill=NAVY; c.alignment=center; c.border=box
ws.row_dimensions[4].height=42
dv_abc=DataValidation(type='list',formula1='"A,B,C"',allow_blank=True,showDropDown=False)
dv_maru=DataValidation(type='list',formula1='"○,△,×"',allow_blank=True,showDropDown=False)
ws.add_data_validation(dv_abc); ws.add_data_validation(dv_maru)
for r in range(5,45):
    for i in range(1,11):
        c=ws.cell(row=r,column=i); c.border=box; c.font=F_BODY
        c.alignment=center if i in (1,2,4,6,8,10) else wrap
        if i==5: c.fill=LBLUE
        if i==8: c.fill=PURPLE_F
        if i==10: c.fill=GREEN_F
    ws.row_dimensions[r].height=30
dv_maru.add('H5:H44'); dv_abc.add('J5:J44')
ws.freeze_panes='D5'
mg(ws,46,1,46,10,'条件②の目安　○＝自分の観点を示し、本文の根拠を引いて主張した　'
                  '△＝発言したが根拠が本文に届かない　×＝発言していない',F_SMALL,GREEN_F,wrapc,h=22)

wb.save(OUT)
print('saved',OUT)
